import os
import io
import time
import json
import shutil
import subprocess
from typing import Optional
import requests
from fastapi import FastAPI, File, UploadFile, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from groq import Groq
from pypdf import PdfReader
from pptx import Presentation

app = FastAPI(title="StudyForge Universal Handler")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

client = Groq(api_key=os.environ.get("GROQ_API_KEY"))


def get_ocr_text(image_path: str) -> str:
    """Sends the extracted frame to OCR.space to avoid CPU load."""
    try:
        with open(image_path, "rb") as f:
            response = requests.post(
                "https://api.ocr.space/parse/image",
                files={image_path: f},
                data={"apikey": "helloworld", "language": "eng"},
                timeout=10,
            )
        result = response.json()
        if not result.get("IsErroredOnProcessing") and result.get("ParsedResults"):
            return result["ParsedResults"][0]["ParsedText"]
    except Exception:
        pass
    return ""


def extract_text_from_pptx(file_bytes: bytes) -> str:
    """Extracts text slide by slide from PowerPoint presentations."""
    prs = Presentation(io.BytesIO(file_bytes))
    text_runs = []
    for i, slide in enumerate(prs.slides):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)
        if slide_texts:
            text_runs.append(f"--- Slide {i + 1} ---\n" + "\n".join(slide_texts))
    return "\n\n".join(text_runs)


def extract_text_from_doc(file_bytes: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))
        pages_text = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages_text).strip()
    elif name.endswith(".pptx"):
        return extract_text_from_pptx(file_bytes)
    elif name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore").strip()
    return ""


@app.post("/generate-quiz")
async def generate_quiz(
    file: UploadFile = File(...),
    num_questions: int = Form(5),
    difficulty: str = Form("Medium"),
    custom_prompt: Optional[str] = Form(None),
):
    start_time = time.time()
    filename = file.filename or "uploaded_file"
    ext = os.path.splitext(filename)[1].lower()

    video_path = f"temp_{filename}"
    audio_path = f"temp_{os.path.splitext(filename)[0]}.mp3"
    image_path = f"temp_{os.path.splitext(filename)[0]}_frame.jpg"

    context_text = ""

    try:
        # 1. Documents (PDF, PPTX, TXT)
        if ext in [".pdf", ".pptx", ".txt"]:
            file_bytes = await file.read()
            context_text = extract_text_from_doc(file_bytes, filename)

        # 2. Audio files (Pure Voice/Lecture recordings)
        elif ext in [".mp3", ".wav", ".m4a", ".aac"]:
            content = await file.read()
            transcription = client.audio.transcriptions.create(
                file=(filename, content),
                model="whisper-large-v3",
                response_format="json",
            )
            context_text = transcription.text

        # 3. Video files (Lectures / Screen records)
        elif ext in [".mp4", ".mkv", ".mov", ".avi", ".webm"]:
            with open(video_path, "wb") as buffer:
                shutil.copyfileobj(file.file, buffer)

            subprocess.run(
                ["ffmpeg", "-y", "-i", video_path, "-vn", "-acodec", "libmp3lame", "-b:a", "64k", audio_path],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                check=True,
            )

            subprocess.run(
                ["ffmpeg", "-y", "-ss", "00:00:03", "-i", video_path, "-frames:v", "1", "-q:v", "2", image_path],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                check=True,
            )

            with open(audio_path, "rb") as af:
                transcription = client.audio.transcriptions.create(
                    file=(audio_path, af.read()),
                    model="whisper-large-v3",
                    response_format="json",
                )
            spoken_text = transcription.text
            slide_text = get_ocr_text(image_path)
            context_text = f"Spoken Transcript:\n{spoken_text}\n\nSlide Content:\n{slide_text}"

        else:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported format '{ext}'. Supported: .pdf, .pptx, .txt, .mp3, .wav, .mp4, .mkv",
            )

        if not context_text or not context_text.strip():
            raise HTTPException(status_code=400, detail="No readable content or transcript found in file.")

        context_preview = context_text[:14000]

        system_instruction = (
            "You are an expert assessment generator. Create multiple-choice questions from the given material. "
            "You MUST respond ONLY with a valid JSON object matching this schema:\n"
            "{\n"
            '  "title": "string",\n'
            '  "notes": "string (summary in markdown)",\n'
            '  "quiz": [\n'
            "    {\n"
            '      "id": 1,\n'
            '      "question": "string",\n'
            '      "options": ["Option A", "Option B", "Option C", "Option D"],\n'
            '      "answer": "Option A",\n'
            '      "explanation": "string"\n'
            "    }\n"
            "  ]\n"
            "}"
        )

        user_content = f"""
Number of Questions: {num_questions}
Target Difficulty: {difficulty}
Custom Instructions: {custom_prompt or "Focus on core concepts and practical understanding."}

Source Content:
---
{context_preview}
---
"""

        completion = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[
                {"role": "system", "content": system_instruction},
                {"role": "user", "content": user_content},
            ],
            response_format={"type": "json_object"},
            temperature=0.3,
            max_tokens=4096,
        )

        result_data = json.loads(completion.choices[0].message.content)
        result_data["processing_time_seconds"] = round(time.time() - start_time, 2)
        return result_data

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    finally:
        for p in [video_path, audio_path, image_path]:
            if os.path.exists(p):
                try:
                    os.remove(p)
                except Exception:
                    pass


@app.get("/")
def root():
    return {"status": "running", "service": "StudyForge Backend"}